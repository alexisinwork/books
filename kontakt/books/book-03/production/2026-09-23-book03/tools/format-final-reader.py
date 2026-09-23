"""Keep short chapter endings together, with optional evidence-led spacing adjustments."""
from pathlib import Path
import argparse,json,hashlib,shutil
from docx import Document
from docx.shared import Pt
p=argparse.ArgumentParser();p.add_argument('--reader',type=Path,required=True);p.add_argument('--spacing',default='{}');p.add_argument('--spacing-file',type=Path);p.add_argument('--page-breaks-file',type=Path);p.add_argument('--archive',type=Path,required=True);a=p.parse_args()
root=Path(__file__).resolve().parents[6]
files=list(a.reader.glob('*.docx'));assert len(files)==1 and not a.archive.exists()
a.archive.mkdir(parents=True)
for f in [files[0],a.reader/'docx-check.json']:
 shutil.copy2(f,a.archive/f.name)
doc=Document(files[0]);original=[x.text for x in doc.paragraphs]
starts=[i for i,x in enumerate(doc.paragraphs) if x.style.name=='Heading 1'];assert len(starts)==32
spacing={int(k):float(v) for k,v in json.loads(a.spacing_file.read_text(encoding='utf-8') if a.spacing_file else a.spacing).items()};changes=[]
for chapter,(start,end) in enumerate(zip(starts,starts[1:]+[len(original)]),1):
 if chapter in spacing:
  for i in range(start+1,end):doc.paragraphs[i].paragraph_format.space_after=Pt(spacing[chapter])
 first=end-1;words=len(original[first].split())
 while first>start+2 and end-first<4 and words+len(original[first-1].split())<=90:
  first-=1;words+=len(original[first].split())
 for i in range(first,end-1):doc.paragraphs[i].paragraph_format.keep_with_next=True
 changes.append({'chapter':chapter,'last_group_paragraphs':end-first,'words':words,'spacing_override_pt':spacing.get(chapter)})
page_breaks=json.loads(a.page_breaks_file.read_text(encoding='utf-8')) if a.page_breaks_file else []
for text in page_breaks:
 matches=[x for x in doc.paragraphs if x.text==text];assert len(matches)==1;matches[0].paragraph_format.page_break_before=True
doc.save(files[0]);assert [x.text for x in Document(files[0]).paragraphs]==original
check=json.loads((a.reader/'docx-check.json').read_text(encoding='utf-8'));check.update(docx_sha256=hashlib.sha256(files[0].read_bytes()).hexdigest(),format_adjustment='Up to four short ending paragraphs kept together; any spacing overrides recorded separately. Prose identity verified.',render='pending')
(a.reader/'docx-check.json').write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
(a.reader/'layout-adjustments.json').write_text(json.dumps({'text_identity_verified':True,'predecessor_archive':a.archive.as_posix(),'chapters':changes,'explicit_page_breaks_before':page_breaks},ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
desktop=Path('C:/Users/alexi/OneDrive/Desktop/KONTAKT-BOOK03-2026-09-23')
for f in [Path(__file__).resolve(),files[0],a.reader/'docx-check.json',a.reader/'layout-adjustments.json',*a.archive.iterdir()]:
 d=desktop/f.resolve().relative_to(root);d.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(f,d)
print(json.dumps({'status':'format_only_text_unchanged','chapters':32}))
