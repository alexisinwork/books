"""Keep short chapter endings together without changing a single prose character."""
from pathlib import Path
import argparse,json,hashlib
from docx import Document
p=argparse.ArgumentParser();p.add_argument('--reader',type=Path,required=True);a=p.parse_args()
files=list(a.reader.glob('*.docx'));assert len(files)==1
doc=Document(files[0]);original=[p.text for p in doc.paragraphs]
starts=[i for i,p in enumerate(doc.paragraphs) if p.style.name=='Heading 1']
assert len(starts)==36
changes=[]
for start,end in zip(starts,starts[1:]+[len(original)]):
    # Keep no more than four final paragraphs and about 90 words together.
    first=end-1;words=len(original[first].split())
    while first>start+2 and end-first<4 and words+len(original[first-1].split())<=90:
        first-=1;words+=len(original[first].split())
    for i in range(first,end-1):doc.paragraphs[i].paragraph_format.keep_with_next=True
    changes.append({'chapter':len(changes)+1,'last_block_paragraphs':end-first,'words':words})
doc.save(files[0]);assert [p.text for p in Document(files[0]).paragraphs]==original
check=json.loads((a.reader/'docx-check.json').read_text(encoding='utf-8'))
check.update(docx_sha256=hashlib.sha256(files[0].read_bytes()).hexdigest(),format_adjustment='Last short group of up to four paragraphs kept together in each chapter; text unchanged.',render='pending')
(a.reader/'docx-check.json').write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
(a.reader/'layout-adjustments.json').write_text(json.dumps({'text_identity_verified':True,'chapters':changes},ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps({'status':'format_only_text_unchanged','chapters':len(changes)}))
