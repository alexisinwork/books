#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    doc = Document()
    for block in args.source.read_text(encoding="utf-8").split("\n\n"):
        text = block.strip()
        if not text:
            continue
        lines = text.splitlines()
        for line in lines:
            if line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            else:
                doc.add_paragraph(line)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(str(args.out).encode("ascii", "backslashreplace").decode("ascii"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
