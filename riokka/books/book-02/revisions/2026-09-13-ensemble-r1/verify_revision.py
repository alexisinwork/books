#!/usr/bin/env python3
"""Verify hashes, synchronized state, reader exports, and DOCX structure."""

from __future__ import annotations

import difflib
from html.parser import HTMLParser
import hashlib
import json
from pathlib import Path
import re

from docx import Document

from finalize_revision import BOOK, BOOK1_SHA, ORIGINAL_SHA, PRIOR, PRIOR_SHA, REV, REVISED, ROOT, printed_count, sha


class Reader(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_main = False
        self.current = None
        self.blocks = []
        self.chapter_ids = []

    def handle_starttag(self, tag, attrs):
        if tag == "main":
            self.in_main = True
        if self.in_main and tag in {"p", "h1", "h2", "h3"}:
            self.current = []
        if self.in_main and tag == "h2":
            self.chapter_ids.append(dict(attrs).get("id"))
        if self.in_main and tag == "br" and self.current is not None:
            self.current.append("\n")

    def handle_data(self, data):
        if self.in_main and self.current is not None:
            self.current.append(data)

    def handle_endtag(self, tag):
        if tag in {"p", "h1", "h2", "h3"} and self.current is not None:
            self.blocks.append("".join(self.current))
            self.current = None
        if tag == "main":
            self.in_main = False


def main() -> None:
    digest = sha(REVISED)
    text = REVISED.read_text(encoding="utf-8")
    assert sha(PRIOR / "revised.md") == PRIOR_SHA
    assert sha(BOOK / "manuscript/master.md") == ORIGINAL_SHA
    assert sha(ROOT / "riokka/books/book-01/manuscript/master.docx") == BOOK1_SHA
    assert printed_count(text) == 400_209
    assert [int(n) for n in re.findall(r"^## Глава (\d+)\.", text, re.M)] == list(range(1, 52))
    assert not re.search(r"\b(?:TODO|FIXME|UNRESOLVED|TBD)\b", text, re.I)

    snapshot = json.loads((REV / "derived/snapshot.json").read_text(encoding="utf-8"))
    assert snapshot["master"]["sha256"] == digest and snapshot["chapter_count"] == 51
    assert "\n".join(p["text"] for p in snapshot["paragraphs"]) + "\n" == text
    assert (REV / "derived/manuscript.txt").read_text(encoding="utf-8") == text
    plain = "\n".join(re.sub(r"^#{1,6} ", "", line).replace("`", "").replace("**", "") for line in text.splitlines()) + "\n"
    assert (REV / "derived/reader.txt").read_text(encoding="utf-8") == plain

    old_text = (PRIOR / "revised.md").read_text(encoding="utf-8")
    expected_diff = "".join(difflib.unified_diff(
        old_text.splitlines(True), text.splitlines(True),
        fromfile="revisions/2026-09-13-10al/revised.md",
        tofile="revisions/2026-09-13-ensemble-r1/revised.md",
    ))
    assert (REV / "changes.diff").read_text(encoding="utf-8") == expected_diff

    chapter_index = json.loads((REV / "derived/chapters.json").read_text(encoding="utf-8"))
    chunks = re.split(r"(?m)(?=^## Глава \d+\.)", text)[1:]
    assert len(chapter_index["chapters"]) == len(chunks) == 51
    for chunk, item in zip(chunks, chapter_index["chapters"]):
        assert hashlib.sha256(chunk.encode()).hexdigest() == item["sha256"]
        assert printed_count(chunk) == item["printed_chars_with_spaces"]

    by_line = {p["p"]: p["text"] for p in snapshot["paragraphs"]}
    b1 = json.loads((ROOT / "riokka/books/book-01/derived/snapshot.json").read_text(encoding="utf-8"))
    b1_lines = {p["p"]: p["text"] for p in b1["paragraphs"]}
    anchors_checked = 0
    book1_anchors_checked = 0

    def walk(value):
        nonlocal anchors_checked, book1_anchors_checked
        if isinstance(value, dict):
            if {"quote", "p", "source_sha256"} <= set(value):
                if value["source_sha256"] == digest:
                    assert by_line[value["p"]].count(value["quote"]) == 1, value
                    anchors_checked += 1
                elif value["source_sha256"] == BOOK1_SHA:
                    assert b1_lines[value["p"]].count(value["quote"]) == 1, value
                    book1_anchors_checked += 1
            for item in value.values():
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)

    json_files = list((REV / "continuity").glob("*.json")) + [REV / "voice.json"]
    for path in json_files:
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["source_sha256"] == digest
        walk(data)
    assert anchors_checked == 282
    assert book1_anchors_checked == 124

    parsed = Reader()
    parsed.feed((REV / "reader.html").read_text(encoding="utf-8"))
    expected_blocks = []
    for block in text.strip().split("\n\n"):
        if re.fullmatch(r"[-*\s]+", block):
            continue
        clean = re.sub(r"^#{1,6} ", "", block).replace("`", "").replace("**", "")
        expected_blocks.append("Конец второго тома." if clean == "*Конец второго тома.*" else clean)
    assert parsed.blocks == expected_blocks
    assert parsed.chapter_ids == [f"chapter-{n}" for n in range(1, 52)]

    book = json.loads((BOOK / "book.json").read_text(encoding="utf-8"))
    assert book["working_revision"]["sha256"] == digest
    assert book["master"]["sha256"] == ORIGINAL_SHA
    for name in ["characters", "end-state", "knowledge", "motifs", "promises", "research", "resources", "scenes", "timeline"]:
        stub = json.loads((BOOK / f"{name}.json").read_text(encoding="utf-8"))
        assert stub["working_revision_view"]["source_sha256"] == digest
    root_issues = json.loads((BOOK / "audit/issues.json").read_text(encoding="utf-8"))
    assert root_issues["working_revision_audit"]["source_sha256"] == digest
    assert root_issues["working_revision_audit"]["open_issue_ids"] == []

    audit = json.loads((REV / "audit/issues.json").read_text(encoding="utf-8"))
    assert not [item for item in audit["items"] if item["resolution_status"] == "unresolved"]
    assert [item["id"] for item in audit["items"] if item["resolution_status"] == "rejected"] == ["ENS-008"]
    metrics = json.loads((REV / "audit/prose-metrics.json").read_text(encoding="utf-8"))
    assert metrics["source_sha256"] == digest and not any(metrics["mechanical_signals"].values())
    sequence = metrics["sequence_checks"]
    assert sequence["lус_before_chapter_21"] == 0
    assert sequence["ray_name_in_garm_chapter_22"] == 0
    assert all(value for key, value in sequence.items() if key not in {"lус_before_chapter_21", "ray_name_in_garm_chapter_22"})

    release = json.loads((REV / "release.json").read_text(encoding="utf-8"))
    assert release["source_sha256"] == digest and release["remaining_decisions"] == []
    for item in release["files"]:
        assert sha(REV / item["file"]) == item["sha256"]
    desktop = Path("/mnt/c/Users/alexi/Desktop/Риокка — Книга 2 — Прочие убытки — ансамблевая редакция для вычитки.docx")
    assert sha(desktop) == release["desktop_delivery"]["docx_sha256"]

    docx_path = REV / "Риокка — Книга 2 — Прочие убытки — ансамблевая редакция.docx"
    doc = Document(docx_path)
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 51
    assert sum(len(p._element.xpath('.//w:br[@w:type="page"]')) for p in doc.paragraphs) == 51

    print(json.dumps({
        "source_sha256": digest,
        "printed_chars_with_spaces": printed_count(text),
        "author_sheets": printed_count(text) / 40_000,
        "chapters": 51,
        "candidate_anchors_checked": anchors_checked,
        "book_1_evidence_anchors_checked": book1_anchors_checked,
        "chapter_hashes": "passed",
        "reader_exports": "passed",
        "continuity_and_metadata": "passed",
        "docx_and_desktop_copy": "passed",
        "unresolved_issues": [],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
