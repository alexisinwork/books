#!/usr/bin/env python3
"""Verify the delivered candidate, its source chain, evidence, and documents."""
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile
from html.parser import HTMLParser
from docx import Document
import fitz
from synchronize import REV, BOOK, ROOT, SOURCE, SOURCE_SHA, BOOK1_SHA, CHANGED, sha, chunks, printed_count, write


class Reader(HTMLParser):
    def __init__(self):
        super().__init__(); self.current=None; self.blocks=[]
    def handle_starttag(self,tag,attrs):
        if tag in ['h1','h2','h3','p']: self.current=[]
        if tag=='br' and self.current is not None: self.current.append('\n')
    def handle_data(self,data):
        if self.current is not None: self.current.append(data)
    def handle_endtag(self,tag):
        if tag in ['h1','h2','h3','p'] and self.current is not None:
            self.blocks.append(''.join(self.current));self.current=None


def main():
    assert sha(SOURCE)==SOURCE_SHA
    original='81641a85ee25e78bc1e57348b279e7a88317551b166cfa831f33bd8f99431d32'
    assert sha(BOOK/'manuscript/master.md')==original
    b1=ROOT/'riokka/books/book-01/manuscript/master.docx'
    assert sha(b1)==BOOK1_SHA
    digest=sha(REV/'revised.md');text=(REV/'revised.md').read_text();lines=text.splitlines()
    before=chunks(SOURCE.read_text());after=chunks(text)
    assert len(before)==len(after)==51
    assert [n for n,(a,b) in enumerate(zip(before,after),1) if a!=b]==CHANGED
    assert re.findall(r'^## Глава (\d+)\.',text,re.M)==[str(n) for n in range(1,52)]
    count=printed_count(text);assert count>=400000
    index=json.loads((REV/'derived/chapters.json').read_text())
    for n,c in enumerate(after): assert hashlib.sha256(c.encode()).hexdigest()==index['chapters'][n]['sha256']
    snapshot=json.loads((REV/'derived/snapshot.json').read_text())
    assert snapshot['master']['sha256']==digest and snapshot['chapter_count']==51
    assert '\n'.join(p['text'] for p in snapshot['paragraphs'])+'\n'==text
    assert (REV/'derived/manuscript.txt').read_text()==text
    expected_plain='\n'.join(re.sub(r'^#{1,6} ','',s).replace('`','').replace('**','') for s in lines)+'\n'
    expected_plain=expected_plain.replace('*Конец второго тома.*','Конец второго тома.')
    assert (REV/'derived/reader.txt').read_text()==expected_plain
    changes=json.loads((REV/'changes.json').read_text());replay=SOURCE.read_text()
    for change in changes['changes']:
        assert replay.count(change['before'])==1
        replay=replay.replace(change['before'],change['after'],1)
    assert replay==text and changes['revised_sha256']==digest

    b1paras=[p.text for p in Document(b1).paragraphs]
    local=external=0
    def evidence(obj):
        nonlocal local,external
        if isinstance(obj,dict):
            if {'quote','p','source_sha256'} <= obj.keys():
                if obj['source_sha256']==digest:
                    assert obj['quote'] in lines[obj['p']-1],obj
                    local+=1
                elif obj['source_sha256']==BOOK1_SHA:
                    assert obj['quote'] in b1paras[obj['p']-1],obj
                    external+=1
                else: raise AssertionError(('unexpected evidence source',obj))
            for v in obj.values():evidence(v)
        elif isinstance(obj,list):
            for v in obj:evidence(v)
    for p in list((REV/'continuity').glob('*.json'))+[REV/'voice.json']:
        data=json.loads(p.read_text());assert data['source_sha256']==digest;evidence(data)
    age=json.loads((REV/'audit/timka-age.json').read_text());evidence(age)
    assert age['age_years']==6 and age['resolution_status']=='resolved'

    patterns={
        'double_space':r' {2,}',
        'space_before_punctuation':r'[ \t]+[,.!?;:]',
        'missing_space_after_comma':r'(?<=[А-Яа-яЁё]),(?=[А-Яа-яЁё])',
        'repeated_word':r'\b([А-Яа-яЁё]{2,})\s+\1\b',
        'editorial_marker':r'\b(?:TODO|FIXME|UNRESOLVED|TBD)\b'}
    raw_signals=[]
    for kind,pattern in patterns.items():
        for m in re.finditer(pattern,text,re.I):
            line=text.count('\n',0,m.start())+1
            known=(kind=='space_before_punctuation' and 'ЖИВЫЕ ДУШИ: 0 из ???' in lines[line-1])
            raw_signals.append(dict(kind=kind,line=line,text=lines[line-1],
                decision='intentional terminal unknown-count placeholder' if known else 'actionable'))
    mechanical={k:sum(x['kind']==k and x['decision']=='actionable' for x in raw_signals) for k in patterns}
    assert not any(mechanical.values()),mechanical
    # Confirm the source's corrected evacuation arithmetic and reveal order.
    assert not re.search(r'\bЛус\w*', ''.join(after[:20]))
    assert 'сто восемьдесят' in after[23] and 'Тридцать одно имя' in after[23]
    assert 'Для Коры оставалось тридцать пять' in after[23]
    assert 'Пять уже принадлежали нам, Веде и обещанному Лусу' in after[23]
    assert 'сто семьдесят шесть' in after[41]
    assert 5*36==180 and 211-180==31 and 35+5==40 and 176+35==211
    assert [after[31].index(t) for t in ['— Шесть минут','— Четыре минуты','— Две минуты','— Минута,','— Успели,']] == sorted(after[31].index(t) for t in ['— Шесть минут','— Четыре минуты','— Две минуты','— Минута,','— Успели,'])

    expected=[]
    for block in text.strip().split('\n\n'):
        if re.fullmatch(r'[-*\s]+',block):continue
        clean=re.sub(r'^#{1,6} ','',block).replace('`','').replace('**','')
        expected.append(clean.replace('*Конец второго тома.*','Конец второго тома.'))
    reader=Reader();reader.feed((REV/'reader.html').read_text());assert reader.blocks==expected
    docx=next(REV.glob('*.docx'));doc=Document(docx)
    actual=[p.text for p in doc.paragraphs if p.text and p.text not in ['Вселенная Риокка','* * *']]
    assert actual==expected
    headings=[p for p in doc.paragraphs if p.style.name=='Heading 1']
    assert len(headings)==51 and all(p.paragraph_format.page_break_before for p in headings)
    with ZipFile(docx) as z,ZipFile(b1) as template:
        for name in ['word/styles.xml','word/settings.xml','word/theme/theme1.xml','word/fontTable.xml']:
            assert z.read(name)==template.read(name),name
        xml=z.read('word/document.xml').decode()
        assert not re.search(r'<w:(?:ins|del|commentReference|commentRangeStart)\b',xml)
    assert [(s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in doc.sections]==[(s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in Document(b1).sections]
    pdf=next((REV/'audit/render').glob('*.pdf'));pages=fitz.open(pdf)
    norm=lambda s:re.sub(r'\s+','',s)
    assert norm('\n'.join(p.get_text(sort=True) for p in pages))==norm('\n'.join(p.text for p in doc.paragraphs))
    render=json.loads((REV/'audit/render/render-check.json').read_text())
    assert render['pdf_sha256']==sha(pdf) and render['docx_sha256']==sha(docx)
    assert render['visual_review_status']=='passed'
    assert not render['blank_pages'] and not render['outside_text_pages']
    lang=json.loads((REV/'audit/language-review.json').read_text())
    assert lang['source_sha256']==digest and lang['unresolved']==0 and lang['matches']==lang['reviewed']

    book=json.loads((BOOK/'book.json').read_text());assert book['working_revision']['sha256']==digest
    assert book['master']['sha256']==original
    for name in ['characters','scenes','timeline','resources','knowledge','promises','motifs','research','end-state','voice']:
        view=json.loads((BOOK/f'{name}.json').read_text())['working_revision_view']
        assert view['source_sha256']==digest and (BOOK/view['path']).is_file()
    release=json.loads((REV/'release.json').read_text());assert release['source_sha256']==digest
    for f in release['files']:assert sha(REV/f['file'])==f['sha256']
    assert sha(Path(release['desktop_delivery']['path']))==sha(docx)
    run=json.loads((REV/'run.json').read_text())
    assert run['status']=='verified'
    assert sha(REV/run['author_decisions']['file'])==run['author_decisions']['sha256']
    assert sha(REV/run['issue_ledger'])==run['issue_ledger_sha256']
    assert sha(REV/run['verification']['file'])==run['verification']['sha256']
    assert sha(REV/run['verification']['report'])==run['verification']['report_sha256']
    ledger=json.loads((REV/'issue-ledger.json').read_text())
    assert all(i['author_decision']['status']=='accepted' and i['verification']['status']=='passed' for i in ledger['items'])
    result=dict(result='passed',file='revised.md',source_sha256=digest,source_before_sha256=SOURCE_SHA,
        printed_chars_with_spaces=count,author_sheets=count/40000,chapters=51,changed_chapters=CHANGED,
        unchanged_chapters=47,local_evidence_anchors=local,book1_evidence_anchors=external,
        mechanical_signals=mechanical,raw_mechanical_findings=raw_signals,docx_sha256=sha(docx),pdf_sha256=sha(pdf),pdf_pages=len(pages),
        docx_pdf_full_text='passed',visual_pages_viewed=len(render['visual_review_pages']),
        language_signals_reviewed=lang['reviewed'],timka_age_years=6,remaining_questions=[])
    write(REV/'audit/verification.json',result)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
