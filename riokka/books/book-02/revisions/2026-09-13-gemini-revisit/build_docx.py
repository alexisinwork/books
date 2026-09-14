#!/usr/bin/env python3
"""Build a Book 2 reading DOCX using Book 1's OOXML package as the template."""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


REV = Path(__file__).resolve().parent
ROOT = REV.parents[4]
TEMPLATE = ROOT / "riokka/books/book-01/manuscript/master.docx"
SOURCE = REV / "revised.md"
OUTPUT = REV / "Риокка — Книга 2 — Прочие убытки — после Gemini.docx"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_font(run, name: str, size: int = 14) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def format_body(paragraph) -> None:
    paragraph.style = "normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_inline(paragraph, text: str) -> None:
    pos = 0
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, "Georgia")
        token = match.group(0)
        run = paragraph.add_run(token[1:-1] if token.startswith("`") else token[2:-2])
        set_font(run, "Courier New" if token.startswith("`") else "Georgia")
        if token.startswith("**"):
            run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, "Georgia")


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document(TEMPLATE)

    # Keep template styles, theme, settings, headers, and A5 section geometry.
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(120)

    title = text.splitlines()[0].removeprefix("# ")
    p = doc.add_paragraph()
    p.style = "normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    set_font(r, "Georgia")
    r.bold = True

    p = doc.add_paragraph()
    p.style = "normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Вселенная Риокка")
    set_font(r, "Georgia")
    r.italic = True

    for block in text.strip().split("\n\n")[1:]:
        if block.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            # A break on an otherwise empty paragraph can itself spill onto
            # a blank page. Bind each chapter's break to its heading instead.
            p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(block[3:])
            set_font(r, "Georgia")
            r.bold = True
        elif block.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(block[4:])
            set_font(r, "Georgia", 13)
            r.bold = True
        elif re.fullmatch(r"[-*\s]+", block):
            p = doc.add_paragraph()
            p.style = "normal"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run("* * *")
            set_font(r, "Georgia")
        elif block == "*Конец второго тома.*":
            p = doc.add_paragraph()
            p.style = "normal"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(24)
            r = p.add_run("Конец второго тома.")
            set_font(r, "Georgia")
            r.italic = True
        else:
            p = doc.add_paragraph()
            format_body(p)
            # A full-line manifest is centered in the Book 1 original.
            if block.startswith("`") and block.endswith("`") and block.count("`") == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
            add_inline(p, block)

    props = doc.core_properties
    props.title = "Контрактник. Книга 2. Прочие убытки"
    props.subject = "Редакция после рекомендаций Gemini для авторской вычитки"
    props.author = ""
    props.comments = "Собрано из проверенного Markdown-кандидата; мастер книги не заменён."
    props.created = datetime(2026, 9, 13, 0, 0, 0, tzinfo=timezone.utc)
    props.modified = datetime(2026, 9, 13, 0, 0, 0, tzinfo=timezone.utc)
    props.revision = 1
    doc.save(OUTPUT)

    # python-docx writes current ZIP timestamps. Normalize them so that a
    # rebuild of the same manuscript has the same SHA-256.
    normalized = OUTPUT.with_suffix(".normalized.docx")
    with ZipFile(OUTPUT) as source_zip, ZipFile(normalized, "w", ZIP_DEFLATED) as target_zip:
        for name in sorted(source_zip.namelist()):
            old = source_zip.getinfo(name)
            info = ZipInfo(name, (2026, 9, 13, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = old.external_attr
            info.create_system = old.create_system
            target_zip.writestr(info, source_zip.read(name))
    normalized.replace(OUTPUT)

    # Structural text check: every Markdown block appears in order after markup removal.
    check = Document(OUTPUT)
    actual = [p.text for p in check.paragraphs if p.text and p.text != "Вселенная Риокка" and p.text != "* * *"]
    expected = []
    for block in text.strip().split("\n\n"):
        if re.fullmatch(r"[-*\s]+", block):
            continue
        clean = re.sub(r"^#{1,6} ", "", block).replace("`", "").replace("**", "")
        if clean == "*Конец второго тома.*":
            clean = "Конец второго тома."
        expected.append(clean)
    if actual != expected:
        raise ValueError("DOCX text extraction does not match the Markdown reading text")
    headings = [p.text for p in check.paragraphs if p.style.name == "Heading 1"]
    if len(headings) != 51:
        raise ValueError(f"expected 51 chapter headings, got {len(headings)}")
    page_breaks = sum(p.paragraph_format.page_break_before is True for p in check.paragraphs)
    if page_breaks != 51:
        raise ValueError(f"expected 51 chapter page breaks, got {page_breaks}")
    print({"output": str(OUTPUT), "sha256": sha(OUTPUT), "chapters": 51, "page_breaks": page_breaks})


if __name__ == "__main__":
    main()
