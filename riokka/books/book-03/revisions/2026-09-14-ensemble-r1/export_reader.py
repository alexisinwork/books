#!/usr/bin/env python3
"""Build the reader DOCX and render it with the prepared private Writer."""
import hashlib
import html
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import fitz
from PIL import Image, ImageDraw, ImageFont

R = Path(__file__).resolve().parent
NAME = 'Риокка — Книга 3 — Холм Розмари — ансамблевая редакция'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def plain(s):
    return re.sub(r'^#{1,6} ', '', s).replace('`', '').replace('*', '').strip()


def write(p, data):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def styled(p, s):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', s):
        if not part:
            continue
        r = p.add_run(part.strip('`*'))
        r.bold = part.startswith('**')
        r.italic = part.startswith('*') and not part.startswith('**')


def build():
    source = R / 'revised.md'
    lines = source.read_text().splitlines()
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.1)
    normal = d.styles['Normal']
    normal.font.name = 'Liberation Serif'
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(3)
    pf.first_line_indent = Cm(.7)
    pf.widow_control = True
    for key, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 15)]:
        s = d.styles[key]
        s.font.name = 'Liberation Serif'
        s.font.size = Pt(size)
        s.font.color.rgb = __import__('docx.shared', fromlist=['RGBColor']).RGBColor(0, 0, 0)
        p = s.paragraph_format
        p.first_line_indent = Cm(0)
        p.space_before = Pt(0)
        p.space_after = Pt(18)
        p.keep_with_next = True
        p.keep_together = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    d.core_properties.title = 'Холм Розмари'
    d.core_properties.subject = 'Контрактник. Книга 3. Мир Риокки'
    d.core_properties.language = 'ru-RU'
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'ru-RU')
    normal.element.get_or_add_rPr().append(lang)
    prev_part = False
    expected = []
    web = ['<!doctype html><html lang="ru"><meta charset="utf-8"><title>Холм Розмари</title>',
           '<style>body{max-width:44em;margin:3em auto;font:19px/1.5 Georgia,serif;padding:0 1em}p{text-indent:1.4em}h1,h2{line-height:1.2}h2{margin-top:3em}@media print{h2{break-before:page}body{font-size:12pt}h1,h2{break-after:avoid}}</style><body>']
    for line in lines:
        if not line.strip() or line.strip() == '---':
            continue
        content = re.sub(r'^#{1,6} ', '', line)
        if line.startswith('# КОНТРАКТНИК.'):
            p = d.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(120)
            styled(p, content)
            prev_part = False
            tag = 'h1'
        elif line.startswith('# ЧАСТЬ '):
            p = d.add_paragraph(style='Heading 1')
            p.paragraph_format.page_break_before = True
            styled(p, content)
            prev_part = True
            tag = 'h1'
        elif line.startswith('## Глава '):
            p = d.add_paragraph(style='Heading 2')
            p.paragraph_format.page_break_before = not prev_part
            styled(p, content)
            prev_part = False
            tag = 'h2'
        else:
            p = d.add_paragraph()
            styled(p, content)
            if line.startswith('*Конец '):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
            tag = 'p'
        expected.append(plain(line))
        web.append(f'<{tag}>{html.escape(plain(line))}</{tag}>')
    web.append('</body></html>')
    target = R / f'{NAME}.docx'
    d.save(target)
    result = [p.text for p in Document(target).paragraphs if p.text.strip()]
    assert result == expected, 'DOCX extraction differs from manuscript'
    (R / 'derived/docx-extracted.txt').write_text('\n'.join(result) + '\n')
    (R / 'derived/reader.txt').write_text('\n'.join(expected) + '\n')
    (R / 'reader.html').write_text('\n'.join(web))
    render = R / 'audit/render'
    render.mkdir(parents=True, exist_ok=True)
    root = Path(json.loads((R / 'renderer-local.json').read_text())['root'])
    binary = root / 'usr/lib/libreoffice/program/soffice'
    assert binary.exists(), 'Run prepare_renderer.py first'
    env = os.environ.copy()
    env['LD_LIBRARY_PATH'] = ':'.join(str(root / x) for x in ['usr/lib/libreoffice/program', 'usr/lib/x86_64-linux-gnu'])
    env['SAL_USE_VCLPLUGIN'] = 'svp'
    env.pop('URE_BOOTSTRAP', None)
    profile = Path(tempfile.mkdtemp(prefix='riokka-book3-ens-writer-'))
    args = [str(binary), f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(render), str(target)]
    proc = subprocess.run(args, env=env, text=True, capture_output=True, timeout=240)
    (render / 'writer.log').write_text(proc.stdout + proc.stderr)
    assert proc.returncode == 0, proc.stderr
    pdf = render / f'{NAME}.pdf'
    assert pdf.exists(), proc.stdout + proc.stderr
    pages = fitz.open(pdf)
    all_text, page_stats, chapter_pages = [], [], {}
    for i, page in enumerate(pages):
        words = page.get_text('words')
        body = [w for w in words if w[3] < page.rect.height - 55]
        # Writer stores the footer before body text. Exclude its fixed region,
        # and check its content separately rather than deleting numeric lines.
        footer_words = [w[4] for w in words if w[1] >= page.rect.height - 55]
        assert footer_words == [str(i + 1)], (i + 1, footer_words)
        text = page.get_text(clip=fitz.Rect(0, 0, page.rect.width, page.rect.height - 55))
        all_text.append(text)
        for m in re.finditer(r'Глава\s+(\d+)\.', text):
            chapter_pages[int(m.group(1))] = i + 1
        bad = [w for w in body if w[0] < 25 or w[2] > page.rect.width - 25 or w[1] < 20]
        page_stats.append({'page': i + 1, 'words': len(body), 'out_of_bounds': len(bad), 'blank_body': not bool(body)})
    import unicodedata
    norm = lambda s: ''.join(c for c in unicodedata.normalize('NFKC', s) if not c.isspace() and c != '\u00ad')
    pdf_text = '\n'.join(all_text)
    extraction_same = norm(pdf_text) == norm('\n'.join(expected))
    (render / 'pdf-extracted.txt').write_text(pdf_text)
    assert extraction_same, 'Rendered PDF text differs from DOCX; inspect before delivery'
    assert len(chapter_pages) == 58
    assert not any(x['out_of_bounds'] or x['blank_body'] for x in page_stats)
    contact_files = []
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 15)
    for start in range(0, len(pages), 20):
        indexes = list(range(start, min(start + 20, len(pages))))
        sheet = Image.new('RGB', (1280, ((len(indexes) + 3) // 4) * 478), '#c4c4c4')
        draw = ImageDraw.Draw(sheet)
        for j, ix in enumerate(indexes):
            pm = pages[ix].get_pixmap(matrix=fitz.Matrix(.52, .52), alpha=False)
            im = Image.frombytes('RGB', [pm.width, pm.height], pm.samples)
            x, y = (j % 4) * 320, (j // 4) * 478
            sheet.paste(im, (x + 5, y + 27))
            draw.text((x + 8, y + 5), str(ix + 1), fill='black', font=font)
        p = render / f'contact-{start + 1:03}-{indexes[-1] + 1:03}.jpg'
        sheet.save(p, quality=86)
        contact_files.append(p.name)
    selections = sorted({1, 2, len(pages), *(chapter_pages[n] for n in [4, 13, 17, 28, 29, 34, 44, 45, 47, 48, 49, 52, 53, 54, 56, 57, 58])})
    for n in selections:
        pages[n - 1].get_pixmap(matrix=fitz.Matrix(1.1, 1.1), alpha=False).save(render / f'page-{n:03}.png')
    write(render / 'verification.json', {'source_sha256': sha(source), 'docx_sha256': sha(target), 'pdf_sha256': sha(pdf),
        'docx_paragraphs': len(expected), 'chapter_count': 58, 'page_count': len(pages), 'docx_text_exact': True,
        'pdf_text_exact_ignoring_layout_whitespace': extraction_same, 'all_page_bounds': page_stats,
        'chapter_pages': chapter_pages, 'contact_sheets': contact_files, 'full_size_samples': selections,
        'visual_review': 'pending_actual_image_inspection'})
    write(R / 'release.json', {'source': 'revised.md', 'source_sha256': sha(source), 'status': 'ensemble_revision_for_author_review',
        'docx': target.name, 'docx_sha256': sha(target), 'pdf': str(pdf.relative_to(R)), 'pdf_sha256': sha(pdf),
        'chapters': 58, 'pages': len(pages), 'reader_text': 'derived/reader.txt', 'verification': 'audit/render/verification.json',
        'desktop_delivery': 'pending', 'github_delivery': 'pending'})
    print(json.dumps({'docx': str(target), 'pages': len(pages), 'paragraphs': len(expected), 'contacts': len(contact_files), 'sha256': sha(source)}, ensure_ascii=False))


if __name__ == '__main__':
    build()
