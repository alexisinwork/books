"""Apply version-bound, exact Russian prose patches while retaining DOCX runs."""
import argparse
from collections import defaultdict
import difflib
import hashlib
import json
from pathlib import Path

from docx import Document
from lxml import etree


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def replace_range(paragraph, start, end, new):
    runs = list(paragraph.runs)
    if ''.join(r.text for r in runs) != paragraph.text:
        raise ValueError('Paragraph contains unsupported text outside direct runs')
    positions, cursor = [], 0
    for run in runs:
        positions.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)
    first = next(i for i, (a, b) in enumerate(positions) if a <= start < b)
    last = next(i for i, (a, b) in enumerate(positions) if a < end <= b)
    prefix = runs[first].text[:start - positions[first][0]]
    suffix = runs[last].text[end - positions[last][0]:]
    if first == last:
        runs[first].text = prefix + new + suffix
    else:
        runs[first].text = prefix + new
        for i in range(first + 1, last):
            runs[i].text = ''
        runs[last].text = suffix


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source', required=True)
    ap.add_argument('--patches', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--reports', required=True)
    args = ap.parse_args()
    source, out, reports = Path(args.source), Path(args.out), Path(args.reports)
    if source.resolve() == out.resolve():
        raise ValueError('Keep the original; output must be a new file')
    data = json.loads(Path(args.patches).read_text(encoding='utf-8'))
    assert sha(source) == data['source_sha256'], 'Source revision mismatch'
    doc = Document(source)
    old_paragraphs = [p.text for p in doc.paragraphs]
    old_xml = [etree.tostring(p._p, method='c14n') for p in doc.paragraphs]
    old_styles = [p.style.name for p in doc.paragraphs]
    original_style_xml = etree.tostring(doc.styles._element, method='c14n')
    groups = defaultdict(list)
    ids = set()
    for patch in data['patches']:
        assert patch['id'] not in ids, 'Repeated patch ID'
        ids.add(patch['id'])
        n, old, new = patch['p'], patch['old'], patch['new']
        assert old and old != new, patch['id']
        text = old_paragraphs[n - 1]
        assert text.count(old) == 1, f"Non-unique old text: {patch['id']} P{n}"
        start = text.index(old)
        groups[n].append({**patch, 'start': start, 'end': start + len(old)})
    for n, patches in groups.items():
        ascending = sorted(patches, key=lambda x: x['start'])
        for left, right in zip(ascending, ascending[1:]):
            assert left['end'] <= right['start'], f"Overlapping patches P{n}: {left['id']} / {right['id']}"
        expected = old_paragraphs[n - 1]
        for patch in reversed(ascending):
            expected = expected[:patch['start']] + patch['new'] + expected[patch['end']:]
            replace_range(doc.paragraphs[n - 1], patch['start'], patch['end'], patch['new'])
        assert doc.paragraphs[n - 1].text == expected
    out.parent.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    check = Document(out)
    new_paragraphs = [p.text for p in check.paragraphs]
    assert len(check.paragraphs) == len(old_paragraphs)
    assert [p.style.name for p in check.paragraphs] == old_styles
    assert etree.tostring(check.styles._element, method='c14n') == original_style_xml
    for i, p in enumerate(check.paragraphs, 1):
        if i not in groups:
            assert p.text == old_paragraphs[i - 1]
            assert etree.tostring(p._p, method='c14n') == old_xml[i - 1], f'Unexpected format change P{i}'
        else:
            assert p.text == doc.paragraphs[i - 1].text
    result = {**data, 'result_sha256': sha(out), 'patch_count': len(ids),
              'changed_paragraphs': len(groups), 'paragraph_count': len(old_paragraphs),
              'status': 'applied_and_text_verified', 'visual_qa': 'pending',
              'unmodified_paragraphs_and_styles_preserved': True}
    (reports / 'changes.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    diff = ''.join(difflib.unified_diff([t+'\n' for t in old_paragraphs], [t+'\n' for t in new_paragraphs],
        fromfile='author-final-before-revision.txt', tofile='editorial-revision.txt', n=3))
    (reports / 'all-changes.diff').write_text(diff, encoding='utf-8')
    lines = ['# Все внесённые правки книги 1', '',
             f"Всего {len(ids)} адресных изменений в {len(groups)} исходных абзацах.", '',
             'Основой был окончательный DOCX, который прочитал брат. Номера P относятся к его 1243 абзацам, включая пустые. В diff каждый исходный абзац — отдельная строка; это стабильнее строк Word, зависящих от шрифта и ширины страницы.', '',
             'SHA-256 исходника: `' + data['source_sha256'] + '`', '',
             'SHA-256 новой редакции: `' + result['result_sha256'] + '`', '',
             'Редакторские решения, затрагивающие новый механизм мира или согласование томов, отдельно перечислены в decisions. Их применение не означает окончательную авторскую приёмку.', '']
    for n in sorted(groups):
        for patch in sorted(groups[n], key=lambda x: x['start']):
            lines += [f"## {patch['id']} — P{n:04d}", '',
                      '**Причина:** ' + patch.get('reason', ''), '',
                      '**Было**', '', *('> ' + x for x in patch['old'].splitlines()), '',
                      '**Стало**', '', *('> ' + x for x in patch['new'].splitlines()), '']
    lines += ['## Решения по замечаниям', '', 'Подробные записи и зависимости сохранены в changes.json.', '']
    decisions = data.get('decisions', [])
    if isinstance(decisions, list):
        for decision in decisions:
            lines += ['- ' + json.dumps(decision, ensure_ascii=False), '']
    else:
        lines += ['```json', json.dumps(decisions, ensure_ascii=False, indent=2), '```']
    (reports / 'ALL-CHANGES.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(json.dumps({k: result[k] for k in ['result_sha256', 'patch_count', 'changed_paragraphs', 'paragraph_count', 'unmodified_paragraphs_and_styles_preserved']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
