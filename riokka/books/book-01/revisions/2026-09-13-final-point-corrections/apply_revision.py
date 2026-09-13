#!/usr/bin/env python3
"""Apply the 19 author-approved point corrections to the updated DOCX."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path

from docx import Document


REPLACEMENTS = [
    (3, "была ли это текущая развилка", "была ли это та самая развилка"),
    (3, "какая во не живет сила", "какая во мне живет сила"),
    (4, "Я - Чистый", "Я — Чистый"),
    (4, "уже сидело то, чего Вешка не смог назвать и от чего отшатнулся: того, кто забирает", "уже поселилось то, чего Вешка не смог назвать и от чего отшатнулся: страх перед тем, кто забирает"),
    (5, "Хотя, я так и не понял чье оружие", "Хотя я так и не понял, чьё оружие"),
    (8, "Раньше и лучше.  —", "Раньше и лучше. —"),
    (8, "меньше,чем тебе", "меньше, чем тебе"),
    (8, "Говорят ты большой фанат", "Говорят, ты большой фанат"),
    (10, "теми же примерно словами", "примерно теми же словами"),
    (11, "в дальнем конце ангара что-то лязгнуло", "в дальнем конце ангара, что-то лязгнуло"),
    (11, "Сорок. Почти вдвое", "Тридцать один. Почти вдвое"),
    (12, "ремесленники с одного цеха", "ремесленники из одного цеха"),
    (15, "И может ты как-нибудь расскажешь под какой", "И, может, ты как-нибудь расскажешь, под какой"),
    (16, "пойти  с сыном", "пойти с сыном"),
    (41, "Я помолчал, и сказал", "Я помолчал и сказал"),
    (41, "унося колонистов, и один бледный росток", "унося колонистов и один бледный росток"),
    (42, "Сати шла первой — медик, который последние дни почти не спал, а теперь распрямился", "Сати шла первой — последние дни она почти не спала, а теперь распрямилась"),
    (42, "несколько местных сортов риокки", "несколько семарийских сортов риокки"),
    (42, "Ворчливый корабль обещал маме рассказать мне про себя еще столько всего", "Мама обещала, что ворчливый корабль расскажет мне про себя ещё столько всего"),
]

CARRYOVER_TEXT_REPLACEMENTS = [
    (
        7,
        "Я глотнул воды, и подумал",
        "Я глотнул воды и подумал",
        "Повторная корректура: однородные сказуемые с одиночным союзом «и».",
    ),
    (
        8,
        "А такой человек, — это уже диагноз",
        "А такой человек — это уже диагноз",
        "Повторная поглавная корректура: перед тире между подлежащим и сказуемым не нужна запятая.",
    ),
    (
        19,
        "Сроки опечатывание Семари",
        "Сроки опечатывания Семари",
        "Исправление формы слова из уже перенесённого критического блока о пятнадцати часах.",
    ),
    (
        19,
        "части системы— как относят",
        "части системы — как относят",
        "Повторная корректура: восстановлены пробелы вокруг тире.",
    ),
    (
        32,
        "Я не нашелся, что ответить",
        "Я не нашелся что ответить",
        "Повторная поглавная корректура: устойчивый оборот «не нашёлся что ответить».",
    ),
    (
        41,
        "Спасибо, брат. Теперь поиграем.»",
        "Спасибо, брат. Теперь поиграем».",
        "Повторная поглавная корректура: точка вынесена за закрывающую кавычку.",
    ),
    (
        41,
        "У него есть ворчливый корабль, который обещала ему мать.",
        "У него есть ворчливый корабль и история про искина, которую обещала ему мать.",
        "Повторная проверка обещаний: Лина обещала историю про искина, а не сам корабль.",
    ),
    (
        42,
        "будто с него сняли мешок",
        "будто с неё сняли мешок",
        "Повторная корректура: местоимение согласовано с Сати после утверждённой перестройки фразы.",
    ),
]

CHAPTER_18_SPLIT = (
    "Пока колонисты собирали первую партию, Джулия разведала еще один путь — старую штольню. "
    "По схеме она была чуть короче и прятала людей от картеля под землёй. Вход и выход Джулия "
    "проверила; в середине сидел охранник, и пройти мимо него она не рискнула. Если наверху станет "
    "жарко, штольня останется единственным ходом. И единственным местом, где нас будут ждать.",
    "К коллектору она шла не пересчитывать зубья капкана. Она искала руку, которая его насторожила: "
    "пружина лжёт хуже человека.",
)

COUNTER_REPLACEMENT = (
    "ЖИВЫЕ ДУШИ: 49 из 50 в секторе.",
    "ЖИВЫЕ ДУШИ: 49 из 50 учтены.",
    "В ПРИЁМНОМ СЕКТОРЕ: 48.",
)

BLANK_PAGE_HEADINGS = {
    "Глава двадцать восьмая. Попутка",
    "Глава тридцать первая. Человек чужой графы",
    "Глава тридцать шестая. Форма без содержания",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_across_runs(paragraph, old: str, new: str) -> list[int]:
    text = paragraph.text
    if text.count(old) != 1:
        raise ValueError(f"Expected exactly one occurrence in paragraph, got {text.count(old)}: {old!r}")
    start = text.index(old)
    end = start + len(old)
    positions = []
    offset = 0
    for index, run in enumerate(paragraph.runs):
        run_end = offset + len(run.text)
        positions.append((index, run, offset, run_end))
        offset = run_end

    affected = [item for item in positions if item[3] > start and item[2] < end]
    if not affected:
        raise ValueError(f"No runs cover replacement: {old!r}")

    first_index, first_run, first_start, _ = affected[0]
    last_index, last_run, last_start, _ = affected[-1]
    prefix = first_run.text[: start - first_start]
    suffix = last_run.text[end - last_start :]

    if first_run is last_run:
        first_run.text = prefix + new + suffix
    else:
        first_run.text = prefix + new
        for _, run, _, _ in affected[1:-1]:
            run.text = ""
        last_run.text = suffix
    return [item[0] for item in affected]


def insert_clone_after(paragraph):
    new_xml = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_xml)
    return paragraph.__class__(new_xml, paragraph._parent)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--log", type=Path, required=True)
    args = parser.parse_args()

    doc = Document(args.source)
    log = []
    for number, (chapter, old, new) in enumerate(REPLACEMENTS, 1):
        old_hits = [i for i, paragraph in enumerate(doc.paragraphs) if old in paragraph.text]
        new_hits = [i for i, paragraph in enumerate(doc.paragraphs) if new in paragraph.text]
        if len(old_hits) != 1 or new_hits:
            raise ValueError(
                f"Replacement {number}: old hits={old_hits}, new hits={new_hits}; expected one old and no new"
            )
        paragraph_index = old_hits[0]
        runs = replace_across_runs(doc.paragraphs[paragraph_index], old, new)
        log.append(
            {
                "id": f"POINT-{number:02d}",
                "chapter": chapter,
                "source_paragraph_index_zero_based": paragraph_index,
                "source_paragraph_label": f"p{paragraph_index + 1}",
                "affected_runs_zero_based": runs,
                "before": old,
                "after": new,
            }
        )

    carryover = []
    for chapter, old, new, reason in CARRYOVER_TEXT_REPLACEMENTS:
        hits = [i for i, paragraph in enumerate(doc.paragraphs) if old in paragraph.text]
        if len(hits) != 1:
            raise ValueError(f"Carryover replacement old hits={hits}; expected exactly one: {old!r}")
        paragraph_index = hits[0]
        runs = replace_across_runs(doc.paragraphs[paragraph_index], old, new)
        carryover.append(
            {
                "action": "replace_text",
                "chapter": chapter,
                "source_paragraph_label": f"p{paragraph_index + 1}",
                "affected_runs_zero_based": runs,
                "before": old,
                "after": new,
                "reason": reason,
            }
        )

    joined = " ".join(CHAPTER_18_SPLIT)
    split_hits = [i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text == joined]
    if len(split_hits) != 1:
        raise ValueError(f"Chapter 18 paragraph split hits={split_hits}; expected exactly one")
    split_index = split_hits[0]
    first_paragraph = doc.paragraphs[split_index]
    second_paragraph = insert_clone_after(first_paragraph)
    first_paragraph.runs[0].text = CHAPTER_18_SPLIT[0]
    second_paragraph.runs[0].text = CHAPTER_18_SPLIT[1]
    carryover.append(
        {
            "action": "split_paragraph",
            "chapter": 18,
            "source_paragraph_label": f"p{split_index + 1}",
            "before": joined,
            "after": list(CHAPTER_18_SPLIT),
            "reason": "Восстановлена утверждённая граница между разведанным путём и чтением ловушки.",
        }
    )

    old_counter, first_counter, second_counter = COUNTER_REPLACEMENT
    counter_hits = [i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text == old_counter]
    if len(counter_hits) != 1:
        raise ValueError(f"Counter paragraph hits={counter_hits}; expected exactly one")
    counter_index = counter_hits[0]
    counter_paragraph = doc.paragraphs[counter_index]
    counter_second = insert_clone_after(counter_paragraph)
    counter_paragraph.runs[0].text = first_counter
    counter_second.runs[0].text = second_counter
    carryover.append(
        {
            "action": "replace_and_split_counter",
            "chapter": 38,
            "source_paragraph_label": f"p{counter_index + 1}",
            "before": old_counter,
            "after": [first_counter, second_counter],
            "reason": "Согласован общий счёт 49 с 48 людьми в секторе и Ореном у рампы.",
        }
    )

    empty_headings = [
        (index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs, 1)
        if paragraph.style.name.lower().startswith(("heading", "заголовок")) and not paragraph.text.strip()
    ]
    if len(empty_headings) != 6:
        raise ValueError(f"Expected six empty heading paragraphs, got {len(empty_headings)}")
    empty_heading_indexes = [index for index, _ in empty_headings]
    for _, paragraph in empty_headings:
        remove_paragraph(paragraph)
    carryover.append(
        {
            "action": "remove_empty_heading_paragraphs",
            "source_paragraph_labels": [f"p{i}" for i in empty_heading_indexes],
            "count": 6,
            "reason": "Удалены шесть пустых абзацев Heading 1 из утверждённой чистки вёрстки.",
        }
    )

    removed_breaks = []
    for heading_text in sorted(BLANK_PAGE_HEADINGS):
        paragraphs = doc.paragraphs
        heading_hits = [i for i, paragraph in enumerate(paragraphs) if paragraph.text == heading_text]
        if len(heading_hits) != 1 or heading_hits[0] == 0:
            raise ValueError(f"Heading hits={heading_hits}; expected one noninitial heading: {heading_text!r}")
        heading_index = heading_hits[0]
        preceding = paragraphs[heading_index - 1]
        if preceding.text.strip() or 'w:type="page"' not in preceding._p.xml:
            raise ValueError(f"Expected a page-break-only paragraph before {heading_text!r}")
        removed_breaks.append(
            {
                "heading": heading_text,
                "source_preceding_paragraph_label_after_prior_cleanup": f"p{heading_index}",
            }
        )
        remove_paragraph(preceding)
    carryover.append(
        {
            "action": "remove_redundant_page_breaks",
            "items": removed_breaks,
            "count": len(removed_breaks),
            "reason": "Удалены три разрыва, создававшие полностью пустые страницы 177, 191 и 220 в контрольном PDF.",
        }
    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    reopened = Document(args.output)
    all_text = "\n".join(p.text for p in reopened.paragraphs)
    for number, (_, old, new) in enumerate(REPLACEMENTS, 1):
        if old in all_text or all_text.count(new) != 1:
            raise ValueError(
                f"Post-save verification failed for replacement {number}: "
                f"old={all_text.count(old)}, new={all_text.count(new)}"
            )
    for _, old, new, _ in CARRYOVER_TEXT_REPLACEMENTS:
        if old in all_text or all_text.count(new) != 1:
            raise ValueError(f"Carryover post-save verification failed: {old!r} -> {new!r}")
    if old_counter in all_text or all_text.count(first_counter) != 1 or all_text.count(second_counter) != 1:
        raise ValueError("Counter post-save verification failed")
    if all_text.count(CHAPTER_18_SPLIT[0]) != 1 or all_text.count(CHAPTER_18_SPLIT[1]) != 1:
        raise ValueError("Chapter 18 split post-save verification failed")
    remaining_empty_headings = [
        paragraph
        for paragraph in reopened.paragraphs
        if paragraph.style.name.lower().startswith(("heading", "заголовок")) and not paragraph.text.strip()
    ]
    if remaining_empty_headings:
        raise ValueError(f"Empty heading paragraphs remain: {len(remaining_empty_headings)}")

    payload = {
        "schema_version": 1,
        "date": "2026-09-13",
        "source": str(args.source),
        "source_sha256": sha256(args.source),
        "output": str(args.output),
        "output_sha256": sha256(args.output),
        "replacement_count": len(log),
        "changes": log,
        "carryover_action_count": len(carryover),
        "carryover_actions": carryover,
    }
    args.log.parent.mkdir(parents=True, exist_ok=True)
    args.log.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
