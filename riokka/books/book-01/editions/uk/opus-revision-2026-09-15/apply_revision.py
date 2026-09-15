#!/usr/bin/env python3
"""Apply patches.json to one base DOCX without touching the base.

Every paragraph of the base has a single run, so text is replaced inside that
run and paragraph/run formatting is preserved. Each patch is guarded: an exact
unique substring, or an `expect` fragment for whole-paragraph operations.
A failed guard aborts the whole build; nothing is written.
"""
import argparse
import hashlib
import json
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--base', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--results', required=True)
    args = ap.parse_args()

    spec = json.loads((HERE / 'patches.json').read_text(encoding='utf-8'))
    doc = Document(args.base)
    paras = list(doc.paragraphs)
    if len(paras) != 1204:
        raise SystemExit(f'Unexpected paragraph count {len(paras)}')
    before = {i: p.text for i, p in enumerate(paras)}
    results = []
    deletions = []

    for patch in spec['patches']:
        p = paras[patch['u']]
        if len(p.runs) != 1:
            raise SystemExit(f"{patch['id']}: paragraph u{patch['u']:04d} has {len(p.runs)} runs")
        run = p.runs[0]
        text = run.text
        op = patch['op']
        status = 'applied'
        if op == 'replace':
            count = text.count(patch['old'])
            if count == 0 and patch.get('only_if_present'):
                status = 'not_applicable_in_base'
            elif count != 1:
                raise SystemExit(f"{patch['id']}: old text found {count} times in u{patch['u']:04d}")
            else:
                run.text = text.replace(patch['old'], patch['new'])
        elif op == 'any_of':
            for old, new in patch['alternatives']:
                if text.count(old) == 1:
                    run.text = text.replace(old, new)
                    break
            else:
                if patch.get('already') and patch['already'] in text:
                    status = 'already_in_base'
                else:
                    raise SystemExit(f"{patch['id']}: no alternative matched u{patch['u']:04d}")
        elif op == 'set':
            if patch['expect'] not in text:
                raise SystemExit(f"{patch['id']}: expect not found in u{patch['u']:04d}")
            run.text = patch['new']
        elif op == 'delete':
            if patch['expect'] not in text:
                raise SystemExit(f"{patch['id']}: expect not found in u{patch['u']:04d}")
            deletions.append(p)
        elif op == 'style':
            if patch['expect'] not in text:
                raise SystemExit(f"{patch['id']}: expect not found in u{patch['u']:04d}")
            old_style = p.style.name
            p.style = doc.styles[patch['style']]
            status = f'applied (style {old_style} -> {patch["style"]})'
        else:
            raise SystemExit(f"{patch['id']}: unknown op {op}")
        results.append({
            'id': patch['id'], 'u': patch['u'], 'op': op, 'severity': patch['severity'],
            'basis': patch['basis'], 'status': status,
            'before': before[patch['u']], 'after': None if op == 'delete' else p.text,
        })

    for p in deletions:
        p._element.getparent().remove(p._element)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    payload = {
        'base': str(args.base), 'base_sha256': digest(args.base),
        'out': str(out), 'out_sha256': digest(out),
        'patches_sha256': digest(HERE / 'patches.json'),
        'paragraphs_before': len(paras), 'paragraphs_after': len(Document(out).paragraphs),
        'results': results,
    }
    Path(args.results).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in payload.items() if k != 'results'}, ensure_ascii=False, indent=2))
    print('statuses:', {s: sum(1 for r in results if r['status'].startswith(s)) for s in {r['status'].split(' ')[0] for r in results}})


if __name__ == '__main__':
    main()
