#!/usr/bin/env python3
"""Build traceable reader documents and mechanical evidence; never overwrite a release."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import language_qa
import uk_naturalness
sys.path.insert(0, str(BASE / 'tools'))
import revise_docx_text as edit

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def string_sha(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()

def save(path, value):
    with path.open('x', encoding='utf-8') as f:
        f.write(value if isinstance(value, str) else json.dumps(value,ensure_ascii=False,indent=2)+'\n')

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--version',required=True)
    args = ap.parse_args()
    out = BASE / args.version
    out.mkdir(exist_ok=False)
    readers = out/'readers'
    readers.mkdir()
    tools = out/'tool-snapshots'
    tools.mkdir()
    for src in (ROOT/'tools/language_qa.py',ROOT/'tools/uk_naturalness.py', language_qa.SNAPSHOT_PATH,
                BASE/'tools/revise_docx_text.py',Path(__file__)):
        shutil.copy2(src,tools/src.name)
    catalog = {'created_utc':datetime.now(timezone.utc).isoformat(),'status':'built_pending_final_gate',
               'author_request':'../REQUEST.md','opus':'explicitly_waived_for_this_cycle',
               'language':'uk','stories':[]}
    for item in json.loads((BASE/'config.json').read_text())['stories']:
        folder = BASE/'stories'/item['id']
        source_manifest = json.loads((folder/'input-manifest.json').read_text())
        for entry in source_manifest['inputs']:
            assert sha(folder/entry['snapshot']) == entry['sha256'], entry
        meta = out/item['id']
        meta.mkdir()
        final = meta/'final.uk.txt'
        shutil.copy2(folder/'uk.txt',final)
        uk = edit.read_projection(final)
        ru = edit.read_projection(folder/'SOURCE/ru.txt')
        desktop = edit.read_projection(folder/'SOURCE/desktop-uk.txt')
        assert len(uk)==len(ru)==len(desktop)==item['expected_units']
        template = folder/'SOURCE/desktop-uk.docx'
        name = item['title']+'.docx'
        docx = readers/name
        template_sha = sha(template)
        before = Document(template)
        assert not before.tables and not before.inline_shapes
        for tag in ('ins','del','hyperlink','fldChar'):
            assert not before._element.findall('.//'+qn('w:'+tag)), f'Unhandled {tag}'
        edit.apply(template,final,docx,template_sha)
        after = Document(docx)
        assert len(before.paragraphs)==len(after.paragraphs)
        for p,q in zip(before.paragraphs,after.paragraphs):
            assert (p._p.pPr.xml if p._p.pPr is not None else None)==(q._p.pPr.xml if q._p.pPr is not None else None)
        for paragraph in after.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    langs = run._r.findall('.//'+qn('w:lang'))
                    assert len(langs)==1 and langs[0].get(qn('w:val'))=='uk-UA'
        assert before.sections[0]._sectPr.xml==after.sections[0]._sectPr.xml
        alignment=[]
        for n,(r,d,u) in enumerate(zip(ru,desktop,uk),1):
            alignment.append({'p':n,'ru_sha256':string_sha(r),'desktop_uk_sha256':string_sha(d),
                              'final_uk_sha256':string_sha(u),'ru_chars':len(edit.plain(r)),
                              'uk_chars':len(edit.plain(u)), 'uk_ru_length_ratio':round(len(edit.plain(u))/max(len(edit.plain(r)),1),3)})
        save(meta/'alignment.json',{'source_sha256':sha(folder/'SOURCE/ru.txt'),'target_sha256':sha(final),
                                    'mapping':'1:1 complete body paragraph mapping; length is only a diagnostic signal','items':alignment})
        changes={}
        for label,old in [('desktop',desktop),('prepared',edit.read_projection(folder/source_manifest['chosen_editorial_base']))]:
            rows=[f'# Було — стало: {item["title"]}\n\nПорівняння з {label}; цільовий SHA-256: `{sha(final)}`.\n']
            count=0
            for n,(a,b) in enumerate(zip(old,uk),1):
                if a!=b:
                    count+=1
                    rows.append(f'\n## Абзац {n}\n\nБуло:\n\n{a}\n\nСтало:\n\n{b}\n')
            changes[label]=count
            save(meta/f'CHANGES-FROM-{label.upper()}.md','\n'.join(rows))
        lang = language_qa.check(final,'uk')
        natural = uk_naturalness.check(final)
        save(meta/'language-qa.json',lang)
        save(meta/'naturalness-qa.json',natural)
        # Prose decisions are recorded by the editor, not inferred from these scans.
        evidence={'story':item['id'],'title':item['title'],'paragraphs':len(uk),'docx':str(docx.relative_to(out)),
                  'docx_sha256':sha(docx),'text_sha256':sha(final),'ru_source_docx_sha256':sha(folder/'SOURCE/ru.docx'),
                  'template_sha256':template_sha,'exact_docx_text':True,'paragraph_properties_preserved':True,
                  'page_geometry_preserved':True,'uk_run_language':True,'changes':changes,
                  'lexical_tokens':lang['lexical_tokens'],'language_signals':len(lang['findings']),
                  'naturalness_signals':len(natural['findings']), 'reading_coverage':'recorded separately; not established by scripts',
                  'source_manifest':str((folder/'input-manifest.json').relative_to(BASE))}
        save(meta/'verification.json',evidence)
        catalog['stories'].append(evidence)
        print(item['id'],changes,'language',len(lang['findings']),'naturalness',len(natural['findings']))
    save(out/'release.json',catalog)

if __name__=='__main__':
    main()
