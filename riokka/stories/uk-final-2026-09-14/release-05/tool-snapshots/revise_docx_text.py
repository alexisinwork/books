#!/usr/bin/env python3
"""Extract or replace DOCX body paragraphs without touching the source file.

The editable projection uses blank lines between Word paragraphs and the
minimal inline markers <b>, </b>, <i>, </i>.  The apply command verifies the
template hash when requested, preserves blank layout paragraphs and paragraph
properties, and checks that the saved DOCX extracts back to the requested text.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree


MARKER = re.compile(r"<(\/)?([bi])>")


def body_paragraphs(doc):
    """Include story paragraphs inside block content controls, in XML order.

    Refuse unhandled tables/text boxes rather than silently lose their text.
    Headers/footers belong to other parts and are not story paragraphs.
    """
    result = []
    allowed = {qn("w:body"), qn("w:sdtContent"), qn("w:sdt")}
    for element in doc._element.body.iter(qn("w:p")):
        parent = element.getparent()
        while parent is not doc._element.body:
            if parent is None or parent.tag not in allowed:
                raise ValueError("Unhandled nested story paragraph")
            parent = parent.getparent()
        result.append(Paragraph(element, doc._body))
    return result


def repair_georgia_embedding(doc):
    """Reader-output-only repair of a mislabeled CJK Georgia embedding.

    This WPS source declares Georgia with Chinese charset 86 and embeds a
    font rendered as SimSun. Keep the requested Georgia family, but remove
    that incorrect embedded substitute. Never modify the frozen source.
    """
    changed = False
    for rel in doc.part.rels.values():
        if not rel.reltype.endswith("/fontTable"):
            continue
        part = rel.target_part
        root = etree.fromstring(part.blob)
        for font in root.findall(qn("w:font")):
            charset = font.find(qn("w:charset"))
            if font.get(qn("w:name")) != "Georgia" or charset is None or charset.get(qn("w:val")) != "86":
                continue
            for child in list(font):
                if etree.QName(child).localname.startswith("embed"):
                    rid = child.get(qn("r:id"))
                    font.remove(child)
                    if rid in part.rels:
                        del part.rels[rid]
                elif child.tag == qn("w:sig"):
                    font.remove(child)
            charset.set(qn("w:val"), "00")
            font.find(qn("w:family")).set(qn("w:val"), "roman")
            font.find(qn("w:pitch")).set(qn("w:val"), "variable")
            changed = True
        if changed:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return changed


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def marked_text(paragraph) -> str:
    chunks: list[str] = []
    bold = italic = False
    for run in paragraph.runs:
        next_bold = run.bold is True
        next_italic = run.italic is True
        if next_bold != bold:
            chunks.append("<b>" if next_bold else "</b>")
            bold = next_bold
        if next_italic != italic:
            chunks.append("<i>" if next_italic else "</i>")
            italic = next_italic
        chunks.append(run.text)
    if italic:
        chunks.append("</i>")
    if bold:
        chunks.append("</b>")
    return "".join(chunks)


def parse_marked(value: str) -> list[tuple[str, bool, bool]]:
    spans: list[tuple[str, bool, bool]] = []
    bold = italic = False
    position = 0
    for marker in MARKER.finditer(value):
        if marker.start() > position:
            spans.append((value[position:marker.start()], bold, italic))
        closing, kind = marker.groups()
        wanted = not bool(closing)
        if kind == "b":
            if bold == wanted:
                raise ValueError(f"Unbalanced bold marker in: {value[:120]!r}")
            bold = wanted
        else:
            if italic == wanted:
                raise ValueError(f"Unbalanced italic marker in: {value[:120]!r}")
            italic = wanted
        position = marker.end()
    if position < len(value):
        spans.append((value[position:], bold, italic))
    if bold or italic:
        raise ValueError(f"Unclosed inline marker in: {value[:120]!r}")
    return spans


def plain(value: str) -> str:
    return MARKER.sub("", value)


def read_projection(path: Path) -> list[str]:
    value = path.read_text(encoding="utf-8")
    if not value.endswith("\n"):
        raise ValueError("Projection must end with a newline")
    paragraphs = value.rstrip("\n").split("\n\n")
    if any(not plain(item).strip() for item in paragraphs):
        raise ValueError("Projection contains an empty editable paragraph")
    for item in paragraphs:
        parse_marked(item)
    return paragraphs


def extract(source: Path, out: Path) -> None:
    if out.exists():
        raise ValueError(f"Output exists: {out}")
    doc = Document(source)
    paragraphs = [marked_text(p) for p in body_paragraphs(doc) if p.text.strip()]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")
    print(f"{out}: {len(paragraphs)} nonempty paragraphs; source_sha256={digest(source)}")


def neutral_base_properties(paragraph):
    runs = [run for run in paragraph.runs if run.text]
    props = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else OxmlElement("w:rPr")
    for tag in ("b", "bCs", "i", "iCs", "lang"):
        for child in props.findall(qn("w:" + tag)):
            props.remove(child)
    # WPS-authored files can carry a CJK document default despite a Western
    # paragraph font. Materialize the intended inherited font, not that fallback.
    if props.find(qn("w:rFonts")) is None:
        style = paragraph.style
        while style is not None:
            if style.font.name:
                fonts = OxmlElement("w:rFonts")
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn("w:" + key), style.font.name)
                props.insert(0, fonts)
                break
            style = style.base_style
    return props


def replace_paragraph(paragraph, marked: str) -> None:
    base = neutral_base_properties(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    for value, bold, italic in parse_marked(marked):
        if not value:
            continue
        run = paragraph.add_run(value)
        run._r.insert(0, deepcopy(base))
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        props = run._r.get_or_add_rPr()
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "uk-UA")
        lang.set(qn("w:eastAsia"), "uk-UA")
        lang.set(qn("w:bidi"), "uk-UA")
        props.append(lang)


def apply(template: Path, projection: Path, out: Path, expected_sha256: str | None) -> None:
    if out.exists():
        raise ValueError(f"Output exists: {out}")
    actual_sha = digest(template)
    if expected_sha256 and actual_sha != expected_sha256:
        raise ValueError(f"Template hash mismatch: expected {expected_sha256}, got {actual_sha}")
    requested = read_projection(projection)
    doc = Document(template)
    editable = [p for p in body_paragraphs(doc) if p.text.strip()]
    if len(editable) != len(requested):
        raise ValueError(f"Paragraph mismatch: DOCX has {len(editable)}, projection has {len(requested)}")
    for paragraph, value in zip(editable, requested):
        replace_paragraph(paragraph, value)
    repair_georgia_embedding(doc)
    doc.core_properties.language = "uk-UA"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    saved = Document(out)
    extracted = [p.text for p in body_paragraphs(saved) if p.text.strip()]
    expected = [plain(p) for p in requested]
    if extracted != expected:
        raise ValueError("Saved DOCX body text differs from the editable projection")
    if digest(template) != actual_sha:
        raise ValueError("Template changed while writing the result")
    print(f"{out}: {len(extracted)} paragraphs; docx_sha256={digest(out)}; template_unchanged=true")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    extract_parser = subparsers.add_parser("extract")
    extract_parser.add_argument("--source", type=Path, required=True)
    extract_parser.add_argument("--out", type=Path, required=True)
    apply_parser = subparsers.add_parser("apply")
    apply_parser.add_argument("--template", type=Path, required=True)
    apply_parser.add_argument("--projection", type=Path, required=True)
    apply_parser.add_argument("--out", type=Path, required=True)
    apply_parser.add_argument("--expected-template-sha256")
    args = parser.parse_args()
    try:
        if args.command == "extract":
            extract(args.source, args.out)
        else:
            apply(args.template, args.projection, args.out, args.expected_template_sha256)
        return 0
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
