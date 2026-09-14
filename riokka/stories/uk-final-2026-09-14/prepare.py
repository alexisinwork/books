#!/usr/bin/env python3
"""Freeze in-scope inputs and create editable projections without changing originals."""
import hashlib
import json
from pathlib import Path
import shutil
import sys
from datetime import datetime, timezone

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
sys.path.insert(0, str(BASE / "tools"))
import revise_docx_text as docx_tool
from docx import Document


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    with path.open("x", encoding="utf-8") as f:
        f.write(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def main():
    for item in json.loads((BASE / "config.json").read_text())["stories"]:
        folder = BASE / "stories" / item["id"]
        folder.mkdir(parents=True, exist_ok=False)
        source = folder / "SOURCE"
        source.mkdir()
        records = []
        for key, name in (("ru", "ru.docx"), ("uk", "desktop-uk.docx")):
            origin = Path(item[key])
            before = sha(origin)
            shutil.copy2(origin, source / name)
            assert sha(source / name) == before == sha(origin)
            doc = Document(source / name)
            assert not doc.tables and not doc.inline_shapes, "Complex DOCX topology requires a separate plan"
            records.append({"role": key, "origin": str(origin), "snapshot": "SOURCE/" + name, "sha256": before,
                            "paragraphs": len(doc.paragraphs), "nonempty": sum(bool(p.text.strip()) for p in doc.paragraphs)})
            docx_tool.extract(source / name, source / ("ru.txt" if key == "ru" else "desktop-uk.txt"))
        chosen = source / "desktop-uk.txt"
        if item["draft"]:
            origin = ROOT / item["draft"]
            before = sha(origin)
            shutil.copy2(origin, source / "prepared-uk.txt")
            assert sha(source / "prepared-uk.txt") == before == sha(origin)
            records.append({"role": "prepared_uk", "origin": str(origin), "snapshot": "SOURCE/prepared-uk.txt", "sha256": before})
            chosen = source / "prepared-uk.txt"
        if item.get("reference_draft"):
            origin = ROOT / item["reference_draft"]
            shutil.copy2(origin, source / "last-delivered-uk.txt")
            records.append({"role": "last_delivered_reference", "origin": str(origin), "snapshot": "SOURCE/last-delivered-uk.txt", "sha256": sha(origin)})
        assert len(docx_tool.read_projection(chosen)) == item["expected_units"]
        shutil.copy2(chosen, folder / "uk.txt")
        save(folder / "input-manifest.json", {"date": datetime.now(timezone.utc).isoformat(), "story": item, "inputs": records,
             "chosen_editorial_base": str(chosen.relative_to(folder)), "selection_basis": "Current author-authorized four-story revision; prepared text used as editable candidate and checked against the frozen RU, not promoted by modification date", "opus": "waived_by_current_author_request"})
        print(item["id"], item["expected_units"], sha(folder / "uk.txt"))


if __name__ == "__main__":
    main()
