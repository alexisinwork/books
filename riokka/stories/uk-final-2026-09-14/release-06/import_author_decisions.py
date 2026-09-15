#!/usr/bin/env python3
"""Import the seven 2026-09-15 author answers into frozen desktop DOCX copies."""
from __future__ import annotations

import hashlib
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent / "tools"))
from revise_docx_text import body_paragraphs, marked_text, plain, replace_paragraph  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

SOURCES = {
    "space": (ROOT / "SOURCE/space-desktop-final.docx", "f2bb885b6a0180286b7d44335c425e0bff3ca7604ece595ffa5497a41d2142f7"),
    "ducks": (ROOT / "SOURCE/ducks-desktop-final.docx", "743b2f8159940a55e9cbf2a743d7b62ebb64604573212bb4947da22ca14c4d10"),
}


def change(value: str, old: str, new: str) -> str:
    count = value.count(old)
    if count != 1:
        raise ValueError(f"Expected one occurrence, found {count}: {old[:90]!r}")
    return value.replace(old, new)


def prepare(kind: str, edits: dict[int, list[tuple[str, str]]], checks: dict[int, str]) -> None:
    source, expected_sha = SOURCES[kind]
    actual_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    if actual_sha != expected_sha:
        raise ValueError(f"Frozen desktop source changed: {source}")
    paragraphs = [marked_text(p) for p in body_paragraphs(Document(source)) if p.text.strip()]
    for number, replacements in edits.items():
        value = paragraphs[number - 1]
        for old, new in replacements:
            value = change(value, old, new)
        paragraphs[number - 1] = value
    for number, needle in checks.items():
        if needle not in plain(paragraphs[number - 1]):
            raise ValueError(f"P{number:03d} missing required final text: {needle!r}")
    projection = ROOT / ("space-is-no-place-for-the-living" if kind == "space" else "where-ducks-fly-in-winter") / "final.uk.txt"
    requested_projection = "\n\n".join(paragraphs) + "\n"
    if projection.exists():
        if projection.read_text(encoding="utf-8") != requested_projection:
            raise ValueError(f"Existing projection differs: {projection}")
    else:
        projection.write_text(requested_projection, encoding="utf-8")
    title = "Космос — не місце для живих" if kind == "space" else "Куди відлітають качки взимку"
    output = ROOT / "readers" / (title + ".docx")
    if output.exists():
        raise ValueError(f"Output exists; archive it before rebuilding: {output}")
    doc = Document(source)
    editable = [p for p in body_paragraphs(doc) if p.text.strip()]
    for number in edits:
        replace_paragraph(editable[number - 1], paragraphs[number - 1])
    if kind == "ducks":
        # Four slightly longer corrected paragraphs would create a page with
        # only the final sentence. Tighten the existing 8 pt gap to 6 pt in
        # the final exchange, without changing margins, fonts or source DOCX.
        for number in range(239, 254):
            editable[number - 1].paragraph_format.space_after = Pt(6)
        all_body = body_paragraphs(doc)
        if len(all_body) != 257 or all_body[-1].text or all_body[-1]._p.getparent() is not doc._element.body:
            raise ValueError("Expected one trailing empty layout paragraph in ducks source")
        doc._element.body.remove(all_body[-1]._p)
    doc.core_properties.language = "uk-UA"
    doc.save(output)
    saved = [p.text for p in body_paragraphs(Document(output)) if p.text.strip()]
    if saved != [plain(v) for v in paragraphs]:
        raise ValueError(f"Saved DOCX text differs from projection: {output}")
    if hashlib.sha256(source.read_bytes()).hexdigest() != expected_sha:
        raise ValueError(f"Source changed while writing: {source}")
    print(f"{output}: {len(saved)} paragraphs; docx_sha256={hashlib.sha256(output.read_bytes()).hexdigest()}; template_unchanged=true")
    print(f"{kind}: {len(paragraphs)} paragraphs, {len(edits)} edited paragraphs")


SPACE_EDITS = {
    97: [("Нам п’ятьом, як старшим за званням, доведеться йти до командного центру ухвалювати стратегічні рішення. І я попрошу тебе взяти з нами Юлю. Ти не проти?",
          "Нам п’ятьом — тобі, мені, Юлі, Таконаші та Шону — як старшим за званням, доведеться йти до командного центру ухвалювати стратегічні рішення. Юля як пілотеса-навігаторка теж потрібна там. Ти не проти?")],
    183: [("Неподалік Марса створили невелику чорну діру.", "Між Юпітером і Сатурном створили невелику чорну діру.")],
    197: [("Саме там, між Сонцем і Землею, вони розмістили станцію: мінімальні витрати палива на підтримання орбіти, цілковита непомітність для нас і повне маскування в разі небезпеки. Судячи з усього, їхня система невидимості не може працювати довго, але щоразу, коли ми пролітали повз, станція була прихована.",
           "Саме там, по той бік Сонця від Землі, вони розмістили станцію: для підтримання орбіти вистачало невеликих регулярних корекцій, із Землі її затуляло Сонце, а від наших кораблів чужинці ховалися маскуванням. Судячи з усього, їхня система невидимості не може працювати довго, але щоразу, коли ми пролітали поблизу, станція була прихована.")],
    234: [("Чорна діра — надмасивний космічний об’єкт, тяжіння якого настільки велике, що вирватися з нього не може навіть світло. Межа цієї області називається горизонтом подій.",
           "Чорна діра — космічний об’єкт, тяжіння якого настільки велике, що вирватися з нього не може навіть світло. Межа цієї області називається горизонтом подій. Чужинська діра — надмасивна за мірками штучних чорних дір: у момент виникнення її маса становила близько ста двадцяти шести тисяч тонн.")],
}

DUCKS_EDITS = {
    2: [("а маю я цілих вісімдесят вісім", "а незабаром мені вісімдесят вісім")],
    3: [("Надворі вже кінець другого тисячоліття", "Надворі вже 2098 рік")],
    6: [("Під навалою спогадів я вирішив узяти дві.", "На полиці стояли пляшечки цього стауту. Під навалою спогадів я вирішив узяти кілька.")],
    17: [("Я взяв свої скромні покупки, порозпихав по кишенях і почимчикував до виходу.",
          "Я взяв свої скромні покупки, склав їх у пакет і почимчикував до виходу.")],
    20: [("Переді мною відкривався бездоганний краєвид: «Височенні дуби й сосни",
          "Переді мною відкривався знайомий краєвид. Я мимоволі згадав його літній вигляд: «Височенні дуби й сосни")],
    85: [("Після зустрічі із сином відсвяткували прийдешнє свято кількома пляшечками ірландського пива.",
          "Після зустрічі із сином відзначили прийдешній день народження, випивши кілька пляшечок ірландського пива.")],
    224: [("У вашому випадку надворі кінець другого тисячоліття.", "У вашому випадку надворі 2098 рік.")],
}


if __name__ == "__main__":
    selected = set(sys.argv[1:]) or {"space", "ducks"}
    if selected - {"space", "ducks"}:
        raise ValueError(f"Unknown story selector: {selected}")
    if "space" in selected:
        prepare("space", SPACE_EDITS, {97: "Юлі, Таконаші та Шону", 183: "Між Юпітером і Сатурном", 197: "по той бік Сонця від Землі", 234: "надмасивна за мірками штучних чорних дір"})
    if "ducks" in selected:
        prepare("ducks", DUCKS_EDITS, {3: "2098 рік", 6: "пляшечки цього стауту", 17: "склав їх у пакет", 20: "згадав його літній вигляд", 85: "кілька пляшечок ірландського пива", 131: "Бога.Нема", 132: "Бога.Нема", 133: "Бога.Нема", 139: "Бога.Нема", 174: "Бога.Нема", 224: "2098 рік"})
