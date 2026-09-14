#!/usr/bin/env python3
import hashlib
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
import sys
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement

BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE/'tools'))
import revise_docx_text as edit
spec=importlib.util.spec_from_file_location('review_runner',BASE.parents[2]/'tools/run_book_review.py')
runner=importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)

class DocumentTests(unittest.TestCase):
    def test_content_control_roundtrip(self):
        with tempfile.TemporaryDirectory(prefix='uk-docx-test-') as temp:
            root=Path(temp)
            doc=Document()
            doc.add_paragraph('Початок')
            nested=doc.add_paragraph('Hidden source')
            control=OxmlElement('w:sdt')
            content=OxmlElement('w:sdtContent')
            control.append(content)
            nested._p.addprevious(control)
            content.append(nested._p)
            doc.add_paragraph('Кінець')
            src=root/'source.docx'
            doc.save(src)
            self.assertEqual(len(Document(src).paragraphs),2)
            self.assertEqual(len(edit.body_paragraphs(Document(src))),3)
            text=root/'target.txt'
            text.write_text('Початок\n\n<b>Українська репліка</b>\n\nКінець\n',encoding='utf-8')
            src_hash=edit.digest(src)
            out=root/'final.docx'
            edit.apply(src,text,out,src_hash)
            actual=[p.text for p in edit.body_paragraphs(Document(out))]
            self.assertEqual(actual,['Початок','Українська репліка','Кінець'])
            self.assertEqual(edit.digest(src),src_hash)
            self.assertTrue(edit.body_paragraphs(Document(out))[1].runs[0].bold)

    def test_table_refused(self):
        doc=Document()
        doc.add_table(rows=1,cols=1).cell(0,0).text='Do not silently omit me'
        with self.assertRaisesRegex(ValueError,'Unhandled nested'):
            edit.body_paragraphs(doc)

    def test_frozen_space_has_two_more_paragraphs(self):
        source=BASE/'stories/space-is-no-place-for-the-living/SOURCE/desktop-uk.docx'
        doc=Document(source)
        self.assertEqual(sum(bool(p.text.strip()) for p in doc.paragraphs),262)
        self.assertEqual(sum(bool(p.text.strip()) for p in edit.body_paragraphs(doc)),264)

    def test_font_repair_is_idempotent(self):
        doc=Document(BASE/'stories/random-experiment/SOURCE/desktop-uk.docx')
        self.assertTrue(edit.repair_georgia_embedding(doc))
        self.assertFalse(edit.repair_georgia_embedding(doc))

class RunnerTests(unittest.TestCase):
    def args(self,root,role='gemini_flash',client='agy'):
        target=root/'target.txt'
        target.write_text('Перша репліка.\n\nДруга репліка.\n',encoding='utf-8')
        prompt=root/'prompt.txt'
        prompt.write_text('Check prose.',encoding='utf-8')
        return SimpleNamespace(out=root/'review',role=role,client=client,model='gemini-test',
                               target=target,source=None,context=[],prompt=prompt)

    def test_stream_transport_and_result(self):
        with tempfile.TemporaryDirectory(prefix='uk-runner-test-') as temp:
            args=self.args(Path(temp))
            digest=runner.sha(args.target)
            report='Status: final\nTarget SHA-256: '+digest+'\n'+('No material defect. '*20)
            response=SimpleNamespace(returncode=0,stdout='\n'.join([
                json.dumps({'event':'init','init':{'model':'gemini-test'}}),
                json.dumps({'event':'result','result':{'status':'SUCCESS','response':report}})]),stderr='')
            with patch.object(runner.subprocess,'run',return_value=response) as mocked:
                runner.run(args)
            command=mocked.call_args.args[0]
            self.assertEqual(command[0],'agy')
            self.assertNotIn('--print',command)
            payload=json.loads(mocked.call_args.kwargs['input'])
            self.assertEqual(payload['event'],'user')
            self.assertIn('[P0002]',payload['message']['content'][0]['text'])
            record=json.loads((args.out/'invocation.json').read_text())
            self.assertEqual(record['status'],'report_returned_pending_manual_validation')
            self.assertEqual(record['target_sha256'],digest)

    def test_wrong_client_rejected(self):
        with tempfile.TemporaryDirectory(prefix='uk-runner-test-') as temp:
            args=self.args(Path(temp),client='claude')
            with patch.object(runner.subprocess,'run') as mocked:
                with self.assertRaisesRegex(ValueError,'only through agy'):
                    runner.run(args)
                mocked.assert_not_called()

    def test_overlong_packet_rejected_before_external_call(self):
        with tempfile.TemporaryDirectory(prefix='uk-runner-test-') as temp:
            args=self.args(Path(temp))
            args.target.write_text('а'*64000+'\n',encoding='utf-8')
            with patch.object(runner.subprocess,'run') as mocked:
                with self.assertRaisesRegex(ValueError,'conservative input bound'):
                    runner.run(args)
                mocked.assert_not_called()

if __name__=='__main__':
    unittest.main()
