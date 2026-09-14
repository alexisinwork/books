#!/usr/bin/env python3
"""Render a reader DOCX and verify its text and every page's basic geometry."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unicodedata

from docx import Document
import fitz
from PIL import Image, ImageDraw, ImageFont
from revise_docx_text import body_paragraphs


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(text):
    return "".join(c for c in unicodedata.normalize("NFKC", text)
                   if not c.isspace() and c != "\u00ad")


def render(source, out, writer_root):
    if out.exists():
        raise ValueError("Choose a new render directory")
    binary = writer_root / "usr/lib/libreoffice/program/soffice"
    if not binary.exists():
        raise FileNotFoundError(binary)
    doc = Document(source)
    for section in doc.sections:
        if any(p.text.strip() for p in section.header.paragraphs + section.footer.paragraphs):
            raise ValueError("Explicit header/footer extraction handling needed")
        if section.header._element.findall('.//' + '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldSimple') or section.footer._element.findall('.//' + '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldSimple'):
            raise ValueError("Page fields need explicit header/footer extraction handling")
    expected = "\n\n".join(p.text for p in body_paragraphs(doc) if p.text.strip()) + "\n"
    out.mkdir(parents=True)
    env = os.environ.copy()
    env["LD_LIBRARY_PATH"] = ":".join(str(writer_root / p) for p in
                                     ["usr/lib/libreoffice/program", "usr/lib/x86_64-linux-gnu"])
    env["SAL_USE_VCLPLUGIN"] = "svp"
    env.pop("URE_BOOTSTRAP", None)
    with tempfile.TemporaryDirectory(prefix="uk-pilot-writer-") as profile:
        command = [str(binary), "-env:UserInstallation=" + Path(profile).as_uri(),
                   "--headless", "--convert-to", "pdf:writer_pdf_Export",
                   "--outdir", str(out.resolve()), str(source.resolve())]
        result = subprocess.run(command, env=env, text=True, capture_output=True, timeout=240)
    (out / "writer.log").write_text(result.stdout + result.stderr)
    pdf = out / (source.stem + ".pdf")
    if result.returncode or not pdf.exists():
        raise RuntimeError("Writer failed; see writer.log")
    pages = fitz.open(pdf)
    extracted, stats = [], []
    margin_notes = []
    # Some source DOCX page-number fields are inside legacy drawing frames,
    # invisible to python-docx's paragraph text. Keep/render them, but compare
    # body text separately. Reject every non-page-number margin item.
    top = doc.sections[0].top_margin.pt
    bottom = doc.sections[0].bottom_margin.pt
    for n, page in enumerate(pages, 1):
        body_rect = fitz.Rect(0, top - 2, page.rect.width, page.rect.height - bottom + 2)
        text = page.get_text(clip=body_rect)
        margin_text = (page.get_text(clip=fitz.Rect(0,0,page.rect.width,top-2))+
                       page.get_text(clip=fitz.Rect(0,page.rect.height-bottom+2,page.rect.width,page.rect.height))).strip()
        if margin_text and normalized(margin_text) != str(n):
            raise ValueError(f"Unexpected margin text on page {n}: {margin_text!r}")
        margin_notes.append({"page":n,"excluded_page_number":margin_text or None})
        extracted.append(text)
        words = page.get_text("words")
        outside = [w for w in words if w[0] < 0 or w[1] < 0 or w[2] > page.rect.width or w[3] > page.rect.height]
        stats.append({"page": n, "words": len(words), "blank": not bool(text.strip()),
                      "out_of_page_bounds": len(outside)})
    text = "\n".join(extracted)
    (out / "pdf-extracted.txt").write_text(text)
    equal = normalized(text) == normalized(expected)
    contacts = []
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 16)
    for start in range(0, len(pages), 16):
        selected = list(range(start, min(start + 16, len(pages))))
        sheet = Image.new("RGB", (1120, ((len(selected) + 3) // 4) * 428), "#bbbbbb")
        draw = ImageDraw.Draw(sheet)
        for index, p in enumerate(selected):
            page = pages[p]
            pix = page.get_pixmap(matrix=fitz.Matrix(265 / page.rect.width, 265 / page.rect.width), alpha=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            x, y = index % 4 * 280 + 7, index // 4 * 428 + 30
            sheet.paste(image, (x, y))
            draw.text((x, y - 24), str(p + 1), font=font, fill="black")
        filename = f"contact-{selected[0]+1:03d}-{selected[-1]+1:03d}.png"
        sheet.save(out / filename)
        contacts.append(filename)
    samples = sorted({1, len(pages), *[n for n, t in enumerate(extracted, 1) if "Розділ" in t]})
    for n in samples:
        pages[n - 1].get_pixmap(matrix=fitz.Matrix(1.15, 1.15), alpha=False).save(out / f"page-{n:03d}.png")
    report = {"source_docx": str(source), "docx_sha256": digest(source), "pdf": pdf.name,
              "pdf_sha256": digest(pdf), "pages": len(pages),
              "pdf_text_exact_ignoring_layout_whitespace": equal, "page_checks": stats,
              "contacts": contacts, "full_size_samples": samples,
              "margin_text_validation": margin_notes,
              "visual_review": "pending_actual_image_inspection", "writer_root": str(writer_root)}
    (out / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    if not equal or any(s["out_of_page_bounds"] for s in stats):
        raise RuntimeError("Rendered text or geometry differs; inspect verification.json")
    print(json.dumps({"pdf": str(pdf), "pages": len(pages), "text_exact": equal, "contacts": contacts}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--writer-root", type=Path, required=True)
    args = parser.parse_args()
    render(args.source, args.out, args.writer_root)
