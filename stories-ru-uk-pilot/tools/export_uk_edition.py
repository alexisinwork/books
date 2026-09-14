#!/usr/bin/env python3
"""Export a checked source-aligned translation into the original DOCX layout."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def paragraph_text(paragraph):
    # Google Docs exports may wrap visible runs in w:sdtContent. python-docx's
    # Paragraph.text omits those runs; source projection includes them.
    tags = {qn("w:t"): None, qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n"}
    return "".join((node.text or "") if tags[node.tag] is None else tags[node.tag]
                   for node in paragraph._p.iter() if node.tag in tags)


def paragraph_runs(paragraph):
    return [Run(node, paragraph) for node in paragraph._p.iter(qn("w:r"))]


def add_text(paragraph, text, base_properties, emphasis):
    spans = []
    for item in emphasis:
        needle = item["text"]
        start = text.find(needle)
        if start < 0 or text.find(needle, start + 1) >= 0:
            raise ValueError("Emphasis must identify one exact target span")
        spans.append((start, start + len(needle), item))
    points = sorted({0, len(text), *(v for a, b, _ in spans for v in (a, b))})
    for start, end in zip(points, points[1:]):
        if end == start:
            continue
        run = paragraph.add_run(text[start:end])
        if base_properties is not None:
            run._r.insert(0, deepcopy(base_properties))
        for a, b, item in spans:
            if a <= start and end <= b:
                if "bold" in item:
                    run.bold = item["bold"]
                if "italic" in item:
                    run.italic = item["italic"]
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "uk-UA")
        props = run._r.get_or_add_rPr()
        for old in props.findall(qn("w:lang")):
            props.remove(old)
        props.append(lang)


def build(source, segments_file, translation_file, out, title):
    if out.exists():
        raise ValueError("Choose a new DOCX output path")
    projection, data = read(segments_file), read(translation_file)
    source_sha = sha(source)
    assert source_sha == data["source_sha256"] == projection["source"]["sha256"]
    assert sha(segments_file) == data["source_segments_sha256"]
    segments = projection["segments"]
    by_id = {item["id"]: item for item in segments}
    expected_ids = [item["id"] for item in segments]
    assert [i for b in data["blocks"] for i in b["source_ids"]] == expected_ids
    doc = Document(source)
    originals = list(doc.paragraphs)
    assert all(paragraph_text(originals[s["p"] - 1]) == s["text"] for s in segments)
    first, covered = {}, set()
    for block in data["blocks"]:
        ids = block["source_ids"]
        first[by_id[ids[0]]["p"]] = block
        covered.update(by_id[i]["p"] for i in ids)
    for p, original in enumerate(originals, 1):
        if paragraph_text(original).strip() and p not in covered:
            raise ValueError(f"Untranslated nonempty source paragraph: {p}")
    expected = [t for b in data["blocks"] for t in b["target_paragraphs"]]
    formatting = []
    for number, original in enumerate(originals, 1):
        if number not in covered:
            continue  # preserve original blank layout paragraphs
        if number not in first:
            original._p.getparent().remove(original._p)
            continue
        block = first[number]
        runs = [r for r in paragraph_runs(original) if r.text.strip()]
        traits = {(r.bold is True, r.italic is True) for r in runs}
        mixed_emphasis = len(traits) > 1 and any(b or i for b, i in traits)
        emphasis = block.get("target_emphasis", [])
        if mixed_emphasis and not emphasis:
            raise ValueError(f"Mixed inline emphasis needs target_emphasis mapping: source p{number}")
        base = deepcopy(runs[0]._r.rPr) if runs else None
        if mixed_emphasis:
            # The first run may itself be emphasized. Keep only traits shared
            # by the whole paragraph as the base; mapped spans restore the rest.
            if base is None:
                base = OxmlElement("w:rPr")
            for trait, position in (("b", 0), ("i", 1)):
                if len({t[position] for t in traits}) > 1:
                    for tag in (trait, trait + "Cs"):
                        for child in base.findall(qn("w:" + tag)):
                            base.remove(child)
                        neutral = OxmlElement("w:" + tag)
                        neutral.set(qn("w:val"), "0")
                        base.append(neutral)
        secondary = [originals[by_id[i]["p"] - 1] for i in block["source_ids"][1:]]
        if any(r.bold or r.italic for p in secondary for r in paragraph_runs(p)) and not emphasis:
            raise ValueError(f"Merged source emphasis needs target_emphasis: source p{number}")
        if any(e.get("paragraph", 0) not in range(len(block["target_paragraphs"])) for e in emphasis):
            raise ValueError(f"Emphasis refers to a missing target paragraph: source p{number}")
        for index, target in enumerate(block["target_paragraphs"]):
            element = deepcopy(original._p)
            for child in list(element):
                if child.tag != qn("w:pPr"):
                    element.remove(child)
            if index and element.pPr is not None:
                for tag in ("pageBreakBefore", "sectPr"):
                    for child in element.pPr.findall(qn("w:" + tag)):
                        element.pPr.remove(child)
            paragraph = Paragraph(element, original._parent)
            spans = [e for e in emphasis if e.get("paragraph", 0) == index]
            add_text(paragraph, target, base, spans)
            original._p.addprevious(element)
        original._p.getparent().remove(original._p)
        if mixed_emphasis or any(b or i for b, i in traits):
            formatting.append({"source_p": number, "mixed": mixed_emphasis,
                               "target_emphasis": emphasis, "uniform_traits": sorted(traits)})
    doc.core_properties.title = title
    doc.core_properties.language = "uk-UA"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    extracted = [p.text for p in Document(out).paragraphs if p.text.strip()]
    assert extracted == expected, "DOCX text differs from the aligned target"
    assert sha(source) == source_sha, "Original DOCX changed"
    report = {"source": str(source), "source_sha256": source_sha,
              "translation": str(translation_file), "translation_sha256": sha(translation_file),
              "docx": out.name, "docx_sha256": sha(out), "paragraphs": len(expected),
              "docx_text_exact": True, "source_unchanged": True,
              "layout": "Original document sections, styles and paragraph formatting retained; target runs replace source text.",
              "source_content_controls": sum(1 for p in originals for _ in p._p.iter(qn("w:sdt"))),
              "emphasis": formatting, "render": "pending"}
    out.with_suffix(".export.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    out.with_suffix(".extracted.txt").write_text("\n\n".join(extracted) + "\n")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    for name in ("source", "segments", "translation", "out"):
        parser.add_argument("--" + name, type=Path, required=True)
    parser.add_argument("--title", required=True)
    args = parser.parse_args()
    build(args.source, args.segments, args.translation, args.out, args.title)
